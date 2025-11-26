from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import requests
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from dhis_utils import dhis_get
from file_utils import resolve_input_path, resolve_output_path


DEFAULT_SHEETS = ("Module 1 - APVD", "Module 2 - APVD")


@dataclass
class ColumnSelector:
    """
    Helper to resolve a column either by index (1-based), Excel letter, or header text.
    """

    raw: str
    index: Optional[int]
    header: Optional[str]

    @classmethod
    def parse(cls, raw: str | int) -> "ColumnSelector":
        text = str(raw).strip()
        if text.isdigit():
            return cls(raw=text, index=int(text), header=None)

        # Try Excel letters (A, B, AA...)
        try:
            numeric_index = column_index_from_string(text)
            return cls(raw=text, index=numeric_index, header=None)
        except ValueError:
            pass

        return cls(raw=text, index=None, header=text)


class UpdateBlueprintDataElementsUseCase:
    """
    Update DHIS2 data element UIDs (and codes when available) inside an XLSX blueprint.
    - Open a workbook.
    - For each configured sheet, read the data element name from a column (by header or position).
    - Ensure the name has the '-APVD' suffix.
    - Look up the data element in DHIS2; if no exact match, ask the user to enter the UID manually.
    - Write UID and code back into the configured columns.
    """

    def __init__(
        self,
        base_url: str,
        jsessionid: str,
        xlsx_path: str,
        output_path: str,
        sheet_names: list[str],
        name_column: str | int,
        uid_column: str | int,
        code_column: str | int,
        data_start_row: int = 2,
        header_scan_rows: int = 5,
    ):
        self.base_url = base_url
        self.jsessionid = jsessionid
        self.xlsx_path = xlsx_path
        self.output_path = output_path
        self.sheet_names = sheet_names
        self.name_selector = ColumnSelector.parse(name_column)
        self.uid_selector = ColumnSelector.parse(uid_column)
        self.code_selector = ColumnSelector.parse(code_column)
        self.data_start_row = data_start_row
        self.header_scan_rows = header_scan_rows

    @staticmethod
    def _ensure_apvd_suffix(name: str) -> str:
        normalized = name.strip()
        if normalized.endswith("-APVD"):
            return normalized
        return f"{normalized}-APVD"

    @staticmethod
    def _normalize_text(value) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _resolve_column(self, sheet, selector: ColumnSelector) -> tuple[int, Optional[int]]:
        """
        Return (column_index, header_row_used).
        If the selector uses a header, search for it in the first header_scan_rows rows.
        """
        if selector.index is not None:
            return selector.index, None

        header_lower = selector.header.lower()
        for row in sheet.iter_rows(min_row=1, max_row=self.header_scan_rows):
            for cell in row:
                value = self._normalize_text(cell.value)
                if value.lower() == header_lower:
                    return cell.column, cell.row

        raise ValueError(
            f"No column with header '{selector.header}' found in first {self.header_scan_rows} rows of sheet '{sheet.title}'"
        )

    def _resolve_columns(self, sheet) -> tuple[dict[str, int], int]:
        """
        Resolve configured columns and compute the first data row (skip header rows if found).
        """
        header_rows: list[int] = []
        columns = {}

        for key, selector in (
            ("name", self.name_selector),
            ("uid", self.uid_selector),
            ("code", self.code_selector),
        ):
            col_index, header_row = self._resolve_column(sheet, selector)
            columns[key] = col_index
            if header_row is not None:
                header_rows.append(header_row)

        data_row_start = self.data_start_row
        if header_rows:
            data_row_start = max(data_row_start, max(header_rows) + 1)

        print(
            f"Sheet '{sheet.title}': name_col={columns['name']}, uid_col={columns['uid']}, "
            f"code_col={columns['code']}, start_row={data_row_start}"
        )
        return columns, data_row_start

    def _prompt_manual_entry(self, sheet_name: str, row_idx: int, target_name: str) -> dict:
        print(
            f"[INPUT REQUIRED] '{target_name}' (sheet '{sheet_name}', row {row_idx}) "
            "could not be matched automatically."
        )
        manual_uid = input("Enter the UID (required, press ENTER to abort): ").strip()
        if not manual_uid:
            raise SystemExit("Aborted by user (no UID provided).")
        manual_code = input("Enter the code (optional, press ENTER to skip): ").strip()
        return {"id": manual_uid, "name": target_name, "code": manual_code or None}

    def _find_data_element(self, target_name: str) -> Optional[dict]:
        """
        Look for a data element whose name matches target_name exactly.
        Returns None if no safe match is found.
        """
        response = dhis_get(
            path="/api/dataElements",
            base_url=self.base_url,
            jsessionid=self.jsessionid,
            params={
                "filter": f"name:like:{target_name}",
                "fields": "id,name,code",
            },
        )

        data_elements = response.get("dataElements") or []
        exact_matches = [
            de for de in data_elements if self._normalize_text(de.get("name")) == target_name
        ]

        if len(exact_matches) == 1:
            return exact_matches[0]

        if len(exact_matches) > 1:
            print(
                f"[WARN] Multiple exact matches for '{target_name}': "
                + ", ".join(de.get("id", "?") for de in exact_matches)
            )
            return None

        if data_elements:
            print(
                f"[WARN] No exact match for '{target_name}'. Candidates: "
                + ", ".join(self._normalize_text(de.get("name")) for de in data_elements)
            )
        else:
            print(f"[WARN] No data elements returned for '{target_name}'.")

        return None

    def _process_row(
        self,
        sheet,
        row_idx: int,
        columns: dict[str, int],
    ) -> bool:
        """
        Process a single row. Returns True if a UID was written.
        """
        name_cell = sheet.cell(row=row_idx, column=columns["name"])
        raw_name = self._normalize_text(name_cell.value)
        if not raw_name:
            return False

        target_name = self._ensure_apvd_suffix(raw_name)
        if target_name != name_cell.value:
            name_cell.value = target_name

        try:
            data_element = self._find_data_element(target_name)
        except requests.HTTPError as http_error:
            status_code = (
                http_error.response.status_code if http_error.response is not None else "?"
            )
            print(
                f"[ERROR] HTTP {status_code} while searching '{target_name}' "
                f"(sheet '{sheet.title}', row {row_idx})"
            )
            data_element = None
        except Exception as unexpected:
            print(
                f"[ERROR] Unexpected error while searching '{target_name}' "
                f"(sheet '{sheet.title}', row {row_idx}): {unexpected}"
            )
            data_element = None

        if data_element is None or self._normalize_text(data_element.get("name")) != target_name:
            data_element = self._prompt_manual_entry(sheet.title, row_idx, target_name)

        uid_value = data_element.get("id")
        if not uid_value:
            data_element = self._prompt_manual_entry(sheet.title, row_idx, target_name)
            uid_value = data_element.get("id")

        sheet.cell(row=row_idx, column=columns["uid"]).value = uid_value

        code_value = data_element.get("code")
        if code_value:
            sheet.cell(row=row_idx, column=columns["code"]).value = code_value

        print(
            f"[OK] Row {row_idx} in '{sheet.title}': "
            f"name='{target_name}', uid='{uid_value}', code='{code_value or 'N/A'}'"
        )
        return True

    def execute(self):
        workbook_path = resolve_input_path(self.xlsx_path)
        if not workbook_path.is_file():
            raise FileNotFoundError(f"Workbook not found at {workbook_path}")

        wb = load_workbook(workbook_path)
        print(f"Opened workbook: {workbook_path}")

        total_updated = 0
        for sheet_name in self.sheet_names:
            if sheet_name not in wb.sheetnames:
                print(f"[WARN] Sheet '{sheet_name}' not found, skipping.")
                continue

            sheet = wb[sheet_name]
            columns, start_row = self._resolve_columns(sheet)

            for row_idx in range(start_row, sheet.max_row + 1):
                if self._process_row(sheet, row_idx, columns):
                    total_updated += 1

        output_path = resolve_output_path(self.output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        print(f"Workbook saved to: {output_path} ({total_updated} rows updated)")
